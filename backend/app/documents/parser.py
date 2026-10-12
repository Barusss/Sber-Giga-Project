from io import BytesIO
from pathlib import Path
import fitz
from docx import Document as WordDocument

class UnsupportedFile(ValueError):
    pass

def extract_pages(data: bytes, extension: str) -> list[str]:
    if extension == ".pdf":
        with fitz.open(stream=data, filetype="pdf") as document:
            pages = []
            for page in document:
                content = page.get_text().strip()
                if not content:
                    # OCR fallback requires Tesseract installed (Docker image includes rus+eng).
                    import pytesseract
                    from PIL import Image
                    pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))
                    picture = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                    content = pytesseract.image_to_string(picture, lang="rus+eng").strip()
                pages.append(content)
            return pages
    if extension == ".docx":
        document = WordDocument(BytesIO(data))
        paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text.strip() for cell in row.cells))
        # DOCX page boundaries cannot reliably be reconstructed without rendering.
        return ["\n".join(paragraphs)]
    raise UnsupportedFile("Supported formats: PDF, DOCX")
